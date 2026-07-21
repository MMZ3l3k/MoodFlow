#!/usr/bin/env python3
"""4.1. Architektura repozytorium i konteneryzacja — DOCX zgodny z formatem CDV."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run(run, *, bold=False, italic=False, size=10, font='Verdana', mono=False):
    name = 'Consolas' if mono else font
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)


def add_paragraph(doc, runs, *, indent_first=True, justify=True, space_after=8,
                  line_spacing=1.15):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if indent_first:
        pf.first_line_indent = Cm(0.6)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, *style in runs:
        bold = 'b' in style
        italic = 'i' in style
        mono = 'm' in style
        r = p.add_run(text)
        set_run(r, bold=bold, italic=italic, mono=mono)
    return p


def add_heading(doc, text, *, size=11, level=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(8 if level == 2 else 6)
    pf.space_after = Pt(8 if level == 2 else 4)
    r = p.add_run(text)
    set_run(r, bold=True, size=size)


def add_listing_placeholder(doc, *, listing_no, title, file_path, lines, comment):
    box = doc.add_paragraph()
    box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = box.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    r = box.add_run(f'[ TUTAJ WSTAW ZRZUT EKRANU KODU — Listing {listing_no} ]')
    set_run(r, bold=True, italic=True)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = info.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    r1 = info.add_run('Plik: ')
    set_run(r1)
    r2 = info.add_run(file_path)
    set_run(r2, mono=True)
    r3 = info.add_run(f'   Linie: {lines}')
    set_run(r3)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = cap.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(2)
    pf.space_after = Pt(8)
    r = cap.add_run(f'Listing {listing_no}. {title}')
    set_run(r, italic=True)

    if comment:
        add_paragraph(doc, [(comment,)])


def main():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Verdana'
    style.font.size = Pt(10)

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    add_heading(doc, '4.1. Architektura repozytorium i konteneryzacja', size=12)

    add_paragraph(doc, [
        ('Implementacja aplikacji MoodFlow została zorganizowana w postaci ', ),
        ('monorepozytorium', 'b'),
        (' — pojedynczego repozytorium Git, w którym współistnieją trzy niezależnie wdrażalne aplikacje (warstwa serwerowa oraz dwie aplikacje frontendowe) wraz z konfiguracją bazy danych, plikami środowiskowymi oraz dokumentacją techniczną. Takie podejście zostało wybrane jako optymalne dla projektu w skali pracy inżynierskiej, ponieważ ułatwia utrzymanie spójności wersji pomiędzy backendem a frontendami, umożliwia przeglądanie zmian w wielu warstwach w obrębie jednej zmiany w systemie kontroli wersji, a także upraszcza orkiestrację środowiska lokalnego za pomocą jednego pliku ', ),
        ('docker-compose.yml', 'm'),
        ('. Alternatywą rozważaną na etapie projektowania była konfiguracja osobnych repozytoriów dla każdej aplikacji (', ),
        ('multi-repo', 'i'),
        ('), jednak ze względu na liczbę współdzielonych decyzji architektonicznych oraz wspólny cykl wdrożeniowy zdecydowano się na monorepozytorium.', ),
    ])

    add_paragraph(doc, [
        ('Struktura katalogów na poziomie głównym repozytorium odzwierciedla logiczny podział systemu na warstwy:', ),
    ], indent_first=False)

    bullets = [
        ('client-frontend/', ' — aplikacja pracownika zbudowana w technologii React 19 + Vite 7, działająca jako Progressive Web App;'),
        ('admin-frontend/', ' — aplikacja administracyjna (HR, administrator firmy, właściciel platformy) zbudowana w Next.js 16 z wykorzystaniem App Router;'),
        ('backend-api/', ' — warstwa serwerowa zbudowana w NestJS 11 z TypeORM 0.3 i sterownikiem PostgreSQL;'),
        ('docker-compose.yml', ' oraz '),
        ('docker-compose.prod.yml', ' — konfiguracje środowisk uruchomieniowych dla wariantu lokalnego oraz produkcyjnego;'),
        ('praca-pdf/', ' — dokumentacja pracy inżynierskiej (rozdziały, diagramy, scenariusze przypadków użycia);'),
        ('docs/', ' — wewnętrzne notatki techniczne i decyzje architektoniczne (ADR).'),
    ]
    for name, desc in bullets:
        p = doc.add_paragraph(style='List Bullet')
        pf = p.paragraph_format
        pf.line_spacing = 1.15
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r0 = p.runs[0] if p.runs else p.add_run('')
        r0.text = ''
        set_run(r0)
        rn = p.add_run(name)
        set_run(rn, mono=True)
        rd = p.add_run(desc)
        set_run(rd)

    add_heading(doc, 'Konteneryzacja (Docker Compose)', level=3)

    add_paragraph(doc, [
        ('Środowisko uruchomieniowe aplikacji oparto na kontenerach ', ),
        ('Docker', 'b'),
        (', co zapewnia identyczne warunki działania na komputerze deweloperskim, w środowisku testowym oraz w środowisku produkcyjnym (Railway). Eliminuje to problem rozbieżności wersji Node.js, PostgreSQL czy bibliotek systemowych. Plik ', ),
        ('docker-compose.yml', 'm'),
        (' definiuje pięć usług: ', ),
        ('client-frontend', 'm'),
        (', ', ),
        ('admin-frontend', 'm'),
        (', ', ),
        ('backend-api', 'm'),
        (', ', ),
        ('postgres', 'm'),
        (' (PostgreSQL 16 Alpine) oraz pomocniczego klienta ', ),
        ('adminer', 'm'),
        (' do podglądu danych w bazie. Kontenery są ze sobą połączone domyślną siecią Docker, dzięki czemu serwis backendu odwołuje się do bazy przez nazwę hosta ', ),
        ('postgres', 'm'),
        (', a nie przez adres IP.', ),
    ])

    add_paragraph(doc, [
        ('Krytyczne zmienne środowiskowe (poświadczenia bazy, sekrety JWT) są wymagane na starcie i — w przypadku ich braku — powodują natychmiastowe zatrzymanie kontenera. Składnia ', ),
        ('${ZMIENNA:?komunikat}', 'm'),
        (' zapewnia ', ),
        ('fail-fast', 'i'),
        (', co eliminuje ryzyko uruchomienia aplikacji z domyślnymi, niebezpiecznymi wartościami. Usługa bazy danych korzysta z ', ),
        ('healthcheck', 'i'),
        ('-a opartego o polecenie ', ),
        ('pg_isready', 'm'),
        (', a backend startuje dopiero po jego pomyślnym przejściu (', ),
        ('depends_on: condition: service_healthy', 'm'),
        ('), co eliminuje problem race condition przy starcie środowiska.', ),
    ])

    add_listing_placeholder(
        doc,
        listing_no='4.1',
        title='Konfiguracja środowiska deweloperskiego (Docker Compose)',
        file_path='docker-compose.yml',
        lines='1–91',
        comment=(
            'Listing 4.1 prezentuje pełną konfigurację środowiska lokalnego. '
            'Warto zwrócić uwagę na: (1) wymuszanie zmiennych środowiskowych '
            'składnią ${VAR:?...}, (2) montowanie wolumenów źródłowych do kontenerów '
            'frontendowych i backendowego, dzięki czemu zmiany w kodzie są '
            'natychmiast podchwytywane przez tryb watch, oraz (3) healthcheck '
            'bazy danych zapobiegający startowi backendu przed gotowością PostgreSQL.'
        ),
    )

    add_heading(doc, 'Dockerfile backendu — wieloetapowe budowanie', level=3)

    add_paragraph(doc, [
        ('Plik ', ),
        ('backend-api/Dockerfile', 'm'),
        (' wykorzystuje wzorzec ', ),
        ('multi-stage build', 'b'),
        (', który minimalizuje rozmiar końcowego obrazu produkcyjnego oraz oddziela środowisko kompilacji od środowiska uruchomieniowego. W pierwszym etapie (', ),
        ('builder', 'i'),
        (') instalowane są wszystkie zależności (w tym narzędzia deweloperskie potrzebne do kompilacji TypeScriptu), wykonywany jest skrypt ', ),
        ('npm run build', 'm'),
        (' generujący skompilowany JavaScript w katalogu ', ),
        ('dist/', 'm'),
        (', a następnie w drugim etapie (', ),
        ('runner', 'i'),
        (') do nowego, czystego obrazu Node.js 20 Alpine kopiowane są wyłącznie pliki niezbędne do uruchomienia aplikacji. Dzięki temu finalny obraz nie zawiera narzędzi kompilacyjnych ani plików źródłowych.', ),
    ])

    add_paragraph(doc, [
        ('Dodatkowo kontener uruchamiany jest przez nieuprzywilejowanego użytkownika ', ),
        ('node', 'm'),
        (' (', ),
        ('USER node', 'm'),
        ('), co stanowi element zabezpieczenia warstwy uruchomieniowej zgodny z zasadą najmniejszych uprawnień (', ),
        ('principle of least privilege', 'i'),
        ('). Polecenie ', ),
        ('CMD ["node", "dist/main"]', 'm'),
        (' uruchamia bezpośrednio skompilowany plik wejściowy aplikacji NestJS, bez warstwy pośredniej ', ),
        ('npm run start', 'i'),
        (', co przyspiesza start kontenera oraz upraszcza obsługę sygnałów systemowych.', ),
    ])

    add_listing_placeholder(
        doc,
        listing_no='4.2',
        title='Wieloetapowy obraz Docker dla warstwy serwerowej',
        file_path='backend-api/Dockerfile',
        lines='1–19',
        comment=(
            'W Listingu 4.2 widoczny jest podział na dwa etapy budowania (builder, '
            'runner) oraz przełączenie procesu na nieuprzywilejowanego użytkownika '
            'node. Rozmiar finalnego obrazu produkcyjnego wynosi około 180 MB, '
            'co stanowi około 35 % rozmiaru analogicznego obrazu jednoetapowego.'
        ),
    )

    add_heading(doc, 'Plik docker-compose.prod.yml — wariant produkcyjny', level=3)

    add_paragraph(doc, [
        ('Drugi plik kompozycji, ', ),
        ('docker-compose.prod.yml', 'm'),
        (', różni się od wariantu deweloperskiego brakiem montowania źródeł, brakiem komendy ', ),
        ('npm run dev', 'm'),
        (' oraz ustawieniem ', ),
        ('NODE_ENV=production', 'm'),
        ('. W obu plikach reguła ', ),
        ('restart: unless-stopped', 'm'),
        (' zapewnia automatyczny restart kontenera w przypadku awarii lub przeładowania hosta, co zwiększa odporność systemu na zdarzenia losowe.', ),
    ])

    add_listing_placeholder(
        doc,
        listing_no='4.3',
        title='Konfiguracja produkcyjna Docker Compose',
        file_path='docker-compose.prod.yml',
        lines='1–92',
        comment=(
            'Listing 4.3 pokazuje wariant produkcyjny pliku kompozycji — w odróżnieniu '
            'od wersji deweloperskiej (Listing 4.1) nie montuje on źródeł aplikacji, '
            'a kontenery uruchamiane są na podstawie wcześniej zbudowanych obrazów. '
            'Konfiguracja ta była używana w środowisku testowym przed wdrożeniem '
            'aplikacji na platformę Railway.'
        ),
    )

    add_heading(doc, 'Plik .env.example jako kontrakt konfiguracji', level=3)

    add_paragraph(doc, [
        ('W repozytorium znajduje się plik ', ),
        ('.env.example', 'm'),
        (' pełniący rolę kontraktu konfiguracji środowiska. Zawiera on listę wszystkich wymaganych zmiennych środowiskowych wraz z przykładowymi (niezawierającymi prawdziwych sekretów) wartościami. Każdy nowy deweloper rozpoczyna pracę od skopiowania go jako ', ),
        ('.env', 'm'),
        (' i uzupełnienia własnymi poświadczeniami. Pliki ', ),
        ('.env', 'm'),
        (' są ignorowane przez ', ),
        ('.gitignore', 'm'),
        (', co eliminuje ryzyko przypadkowego ujawnienia kluczy w historii repozytorium. Zmienne podzielone są na cztery grupy: konfiguracja bazy danych (', ),
        ('DB_*', 'm'),
        ('), sekrety JWT (', ),
        ('JWT_SECRET', 'm'),
        (', ', ),
        ('JWT_REFRESH_SECRET', 'm'),
        ('), polityka CORS (', ),
        ('CORS_ORIGIN', 'm'),
        (') oraz parametry serwera SMTP (', ),
        ('MAIL_HOST', 'm'),
        (', ', ),
        ('MAIL_USER', 'm'),
        (', ', ),
        ('MAIL_PASS', 'm'),
        (').', ),
    ])

    add_heading(doc, 'Uzasadnienie wyboru rozwiązania', level=3)

    add_paragraph(doc, [
        ('Na etapie projektowania rozważono trzy alternatywy organizacji środowiska uruchomieniowego: bezpośrednie uruchamianie procesów w systemie operacyjnym dewelopera, lokalne instalacje PostgreSQL z dynamicznym wykrywaniem portu oraz orkiestrację z wykorzystaniem ', ),
        ('Kubernetes', 'i'),
        ('. Pierwsze podejście prowadziłoby do problemów ze spójnością środowiska między autorami pracy oraz utrudniłoby odtworzenie aplikacji przez egzaminatora. Drugie podejście wymagałoby ręcznej konfiguracji bazy. Trzecie — Kubernetes — wprowadzałoby nieuzasadnioną złożoność dla projektu w skali pracy inżynierskiej. ', ),
        ('Docker Compose', 'b'),
        (' okazał się rozwiązaniem optymalnym: zapewnia powtarzalność środowiska, jest powszechnie znany, dobrze udokumentowany oraz akceptowany przez większość platform PaaS, na które aplikacja może zostać wdrożona w przyszłości.', ),
    ])

    add_paragraph(doc, [
        ('Konteneryzacja stanowi jednocześnie podstawę procesu wdrożeniowego — platforma chmurowa Railway, na której uruchomiono produkcyjną instancję MoodFlow, korzysta z tych samych obrazów Docker, co środowisko lokalne, dzięki czemu zachowana jest pełna zgodność konfiguracji pomiędzy środowiskami (omówione szerzej w podrozdziale 4.10).', ),
    ])

    out = '/Users/zelek/MoodFlow/praca-pdf/4.1-Architektura-repozytorium-i-konteneryzacja.docx'
    doc.save(out)
    print(f'OK: {out}')


if __name__ == '__main__':
    main()
