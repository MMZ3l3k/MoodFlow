#!/usr/bin/env python3
"""
Zwięzły dokument:
  Metodologia pracy + Wykorzystywane narzędzia + Stack technologiczny

Krótka forma, głównie tabele i wypunktowania.

Wyjście: Metodologia-Narzedzia-Stack.docx
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


NAVY      = "1D3E44"
HEADER_BG = "1D3E44"
ALT_BG    = "F4F1E8"


# ---------- helpers ----------
def set_run(run, *, bold=False, italic=False, size=10, color=None, font="Verdana"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:cs"), font)


def add_heading(doc, text, *, size=14, color=NAVY, space_before=14, space_after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    r = p.add_run(text)
    set_run(r, bold=True, size=size, color=color)


def add_subheading(doc, text, *, size=11, color=NAVY, space_before=8, space_after=4):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    r = p.add_run(text)
    set_run(r, bold=True, size=size, color=color)


def add_para(doc, runs, *, indent_first=False, justify=True, space_after=4, size=10):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if indent_first:
        pf.first_line_indent = Cm(0.6)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if isinstance(runs, str):
        r = p.add_run(runs); set_run(r, size=size)
    else:
        for item in runs:
            if isinstance(item, str):
                r = p.add_run(item); set_run(r, size=size)
            else:
                text, *style = item
                r = p.add_run(text)
                set_run(r, bold="b" in style, italic="i" in style, size=size)
    return p


def add_bullet(doc, content, *, size=10):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.line_spacing = 1.2
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if p.runs:
        p.runs[0].text = ""
        set_run(p.runs[0], size=size)
    if isinstance(content, str):
        r = p.add_run(content); set_run(r, size=size)
    else:
        for item in content:
            if isinstance(item, str):
                r = p.add_run(item); set_run(r, size=size)
            else:
                text, *style = item
                r = p.add_run(text)
                set_run(r, bold="b" in style, italic="i" in style, size=size)


# ---- table cell helpers ----
def shade_cell(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    for s in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(s)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def set_cell_border(cell, color="999999", size=4):
    tc_pr = cell._tc.get_or_add_tcPr()
    for tb in tc_pr.findall(qn("w:tcBorders")):
        tc_pr.remove(tb)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:color"), color)
        borders.append(b)
    tc_pr.append(borders)


def set_cell_width(cell, w_cm):
    cell.width = Cm(w_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    for tcW in tc_pr.findall(qn("w:tcW")):
        tc_pr.remove(tcW)
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:type"), "dxa")
    tcW.set(qn("w:w"), str(int(w_cm * 567)))
    tc_pr.append(tcW)


def set_table_grid(table, widths_cm):
    tbl = table._element
    tbl_pr = tbl.find(qn("w:tblPr"))
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        tbl.remove(tblGrid)
    tblGrid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w * 567)))
        tblGrid.append(gc)
    tbl_pr.addnext(tblGrid)


def write_cell(cell, text, *, bold=False, italic=False, color=None, size=9,
               bg=None, align="left", anchor="center"):
    cell.text = ""
    cell.vertical_alignment = {
        "top": WD_ALIGN_VERTICAL.TOP,
        "center": WD_ALIGN_VERTICAL.CENTER,
        "bottom": WD_ALIGN_VERTICAL.BOTTOM,
    }[anchor]
    p = cell.paragraphs[0]
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(1); pf.space_after = Pt(1)
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    set_cell_border(cell)
    if bg:
        shade_cell(cell, bg)
    if text:
        r = p.add_run(text)
        set_run(r, bold=bold, italic=italic, color=color, size=size)


def add_table(doc, rows, col_widths):
    n_cols = len(col_widths)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, col_widths)
    for r_idx, row in enumerate(rows):
        is_header = (r_idx == 0)
        for c_idx, val in enumerate(row):
            bg = HEADER_BG if is_header else (ALT_BG if r_idx % 2 == 0 else "FFFFFF")
            color = "FFFFFF" if is_header else None
            write_cell(
                table.cell(r_idx, c_idx), str(val),
                bold=is_header,
                size=9 if not is_header else 10,
                color=color,
                bg=bg,
                align="center" if is_header else "left",
            )
            set_cell_width(table.cell(r_idx, c_idx), col_widths[c_idx])
    return table


# =========================================================
# BUDOWA DOKUMENTU
# =========================================================
def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Verdana"
    style.font.size = Pt(10)

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width  = Cm(21.0)
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.2)
    section.right_margin  = Cm(1.8)

    # tytuł na początku
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Metodologia, narzędzia i stack technologiczny — MoodFlow")
    set_run(r, bold=True, size=16, color=NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Krótkie zestawienie sposobu pracy zespołu, używanego oprzyrządowania i technologii wybranych do realizacji projektu.")
    set_run(r, italic=True, size=10, color="666666")

    # =========================================================
    # 1. METODOLOGIA PRACY
    # =========================================================
    add_heading(doc, "1.  Metodologia pracy", size=14, space_before=4)

    add_subheading(doc, "Model pracy")
    for item in [
        [("Scrum-lite + Kanban", "b"), " — hybryda zwinna dostosowana do dwuosobowego zespołu studenckiego."],
        [("Iteracje 2-tygodniowe", "b"), " — zamknięty zakres funkcjonalności na sprint, krótki planning + retro."],
        [("Asynchroniczna komunikacja", "b"), " — codzienne updaty statusu zastępują formalne stand-upy."],
        [("Praca w parze przy modułach krytycznych", "b"), " — auth, multi-tenant, anonimizacja, scoring."],
    ]:
        add_bullet(doc, item)

    add_subheading(doc, "Tablica zadań (Kanban)")
    for item in [
        [("Backlog", "b"), " — pomysły i zadania bez przypisanej iteracji."],
        [("To do", "b"), " — zadania zaplanowane na bieżący sprint."],
        [("In progress", "b"), " — zadania w trakcie realizacji (limit WIP = 2 / osobę)."],
        [("Done", "b"), " — zadania zakończone i zatwierdzone do scalenia."],
    ]:
        add_bullet(doc, item)

    add_subheading(doc, "Priorytetyzacja zakresu — MoSCoW")
    add_para(doc, [
        "Każda funkcjonalność z PRD ma przypisany jeden z czterech priorytetów: ",
        ("Must Have", "b"), " (zakres MVP, krytyczne dla obrony), ",
        ("Should Have", "b"), " (istotne, planowane do MVP), ",
        ("Could Have", "b"), " (jeśli zostanie czas), ",
        ("Won't Have", "b"), " (poza zakresem pracy). Klasyfikacja jest źródłem prawdy dla harmonogramu (rozdział 4) i decyzji o zakresie.",
    ])

    add_subheading(doc, "Definition of Done")
    for item in [
        "kod przeglądnięty w PR przez drugiego członka zespołu;",
        "ESLint i typecheck TypeScript nie zgłaszają błędów;",
        "endpoint backendowy pokryty testem klikalnym lub jednostkowym;",
        "frontend manualnie sprawdzony w obu aplikacjach (pracownik + admin);",
        "zmiana udokumentowana w README, komentarzu PR lub rozdziale pracy.",
    ]:
        add_bullet(doc, item)

    add_subheading(doc, "Kontrola wersji")
    for item in [
        [("Monorepozytorium Git", "b"), " z trzema strefami: backend-api, client-frontend, admin-frontend."],
        [("Pull Requesty", "b"), " — jedyna droga zmian na master; obowiązkowy review."],
        [("Conventional Commits (PL)", "b"), " — np. ", ("fix(client-frontend): wymuś jasne tokeny tekstu na ekranach auth", "i"), "."],
        [("Migracje TypeORM", "b"), " — każda zmiana schematu posiada plik forward + revert."],
    ]:
        add_bullet(doc, item)

    doc.add_page_break()

    # =========================================================
    # 2. STACK TECHNOLOGICZNY
    # =========================================================
    add_heading(doc, "2.  Stack technologiczny", size=14, space_before=0)

    add_para(doc, "Aplikacja webowa trójwarstwowa: dwie aplikacje frontendowe (pracownik / HR + admin) komunikują się z centralnym API. Wszystkie warstwy w TypeScript.")

    add_subheading(doc, "Backend (backend-api)")
    rows = [
        ["Element", "Technologia", "Uzasadnienie"],
        ["Framework",        "NestJS 11",                 "modularna struktura, DI, podział na kontrolery/serwisy"],
        ["Język",            "TypeScript 5.9",            "statyczne typowanie DTO i modeli domenowych"],
        ["ORM",              "TypeORM 0.3",                "migracje + repozytoria, dobra integracja z NestJS"],
        ["Baza danych",      "PostgreSQL 16",              "relacyjna, wielotabelowa: użytkownicy + ankiety + wyniki"],
        ["Autentykacja",     "JWT + Passport",             "access + refresh, RBAC pracownik/HR/admin"],
        ["Hashowanie haseł", "bcrypt",                     "standard branżowy"],
        ["Walidacja DTO",    "class-validator",            "deklaratywne reguły na poziomie kontrolerów"],
        ["Rate limiting",    "@nestjs/throttler",          "ochrona logowania przed brute-force"],
        ["Nagłówki HTTP",    "helmet",                     "podstawowa ochrona XSS / clickjacking"],
        ["E-mail",           "Nodemailer",                 "powiadomienia (zatwierdzenie konta, nowy test)"],
    ]
    add_table(doc, rows, col_widths=[3.5, 3.5, 9.5])

    add_subheading(doc, "Frontend pracownika (client-frontend)")
    rows = [
        ["Element", "Technologia", "Uzasadnienie"],
        ["Framework UI",  "React 19",            "komponentowy model UI"],
        ["Build tool",    "Vite 8",              "szybki HMR, krótkie czasy buildu"],
        ["Style",         "Tailwind CSS 3",      "utility-first, design system w klasach"],
        ["State",         "Redux Toolkit 2",     "store sesji i meta-danych użytkownika"],
        ["Routing",       "React Router 7",      "ochrona widoków przez strażników"],
        ["Animacje",      "Framer Motion 12",    "subtelne przejścia"],
        ["Ikony",         "Lucide React",        "spójna paleta line-art"],
        ["PWA",           "vite-plugin-pwa",     "instalowalność + offline po pierwszym wejściu"],
    ]
    add_table(doc, rows, col_widths=[3.5, 3.5, 9.5])

    doc.add_page_break()

    add_subheading(doc, "Frontend administracyjny / HR (admin-frontend)", space_before=0)
    rows = [
        ["Element", "Technologia", "Uzasadnienie"],
        ["Framework UI",  "Next.js 16 (App Router)", "uporządkowana struktura panelu"],
        ["Style",         "Tailwind CSS 4",          "ten sam design system co u pracownika"],
        ["Wykresy",       "Recharts 3",              "trendy HR (linowe i słupkowe)"],
        ["Eksport",       "jsPDF + html2canvas",     "raport HR do PDF po stronie klienta"],
        ["Ciasteczka",    "js-cookie",               "zarządzanie sesją w panelu"],
        ["Theming",       "next-themes",             "tryb jasny / ciemny"],
        ["PWA",           "@ducanh2912/next-pwa",    "instalowalność panelu HR"],
    ]
    add_table(doc, rows, col_widths=[3.5, 3.5, 9.5])

    add_subheading(doc, "Infrastruktura i wdrożenie")
    rows = [
        ["Element", "Technologia", "Uzasadnienie"],
        ["Konteneryzacja",   "Docker + Docker Compose", "powtarzalne środowisko dev i prod"],
        ["Środowisko prod",  "Railway",                "orkiestracja kontenerów + managed PostgreSQL"],
        ["VCS",              "Git + GitHub",           "monorepozytorium, ochrona master, PR-y"],
        ["CI (planowane)",   "GitHub Actions",         "lint + typecheck + testy klikalne przy każdym PR"],
    ]
    add_table(doc, rows, col_widths=[3.5, 3.5, 9.5])

    doc.add_page_break()

    # =========================================================
    # 3. WYKORZYSTYWANE NARZĘDZIA
    # =========================================================
    add_heading(doc, "3.  Wykorzystywane narzędzia", size=14, space_before=0)

    add_para(doc, "Zestaw narzędzi został dobrany tak, by zminimalizować tarcia pomiędzy pracą developerską a pracą dokumentacyjną. Większość pozycji jest darmowa w wariancie studenckim lub open-source.")

    add_subheading(doc, "Narzędzia developerskie")
    rows = [
        ["Kategoria", "Narzędzie", "Zastosowanie"],
        ["IDE",                   "Visual Studio Code, Cursor",      "edycja kodu + integracja z AI"],
        ["Środowisko",            "Docker Desktop",                  "baza i usługi lokalnie"],
        ["Klient bazy danych",    "TablePlus, pgAdmin",              "przeglądanie schematu, ad-hoc query"],
        ["Klient HTTP",           "Postman, REST Client (VS Code)",  "testowanie endpointów API"],
        ["Lint / formatter",      "ESLint, Prettier",                "konwencje kodu"],
        ["Typecheck",             "TypeScript Compiler",             "weryfikacja typów w obu warstwach"],
        ["Migracje DB",           "TypeORM CLI",                     "wersjonowanie schematu"],
        ["Build",                 "Vite, Next.js build, nest build", "produkcyjne paczki"],
        ["Testy",                 "Jest (backend), Vitest (front)",  "testy jednostkowe i integracyjne"],
    ]
    add_table(doc, rows, col_widths=[3.5, 4.5, 8.5])

    add_subheading(doc, "Narzędzia współpracy i dokumentacji")
    rows = [
        ["Kategoria", "Narzędzie", "Zastosowanie"],
        ["Kontrola wersji",    "GitHub",                       "PR-y, code review, ochrona master"],
        ["Tablica zadań",      "GitHub Projects (Kanban)",     "śledzenie sprintów i statusu zadań"],
        ["Komunikacja",        "Discord, Slack",               "krótkie ustalenia i wymiana plików"],
        ["Spotkania",          "Google Meet, MS Teams",        "konsultacje z promotorem, sync zespołu"],
        ["Diagramy",           "draw.io, Mermaid, PlantUML",   "ERD, przypadki użycia, sekwencje"],
        ["Design UI",          "Figma, Pencil",                "makiety i prototypy, biblioteka komponentów"],
        ["Pisanie pracy",      "Microsoft Word, python-docx",  "rozdziały + generowanie tabel"],
        ["Notatki / wiedza",   "Notion, Obsidian",             "notatki badawcze, ADR"],
        ["Wsparcie AI",        "Claude Code, ChatGPT",         "review kodu, generowanie fragmentów docs"],
        ["Hasła i sekrety",    "1Password",                    "bezpieczne przechowywanie sekretów rozwojowych"],
    ]
    add_table(doc, rows, col_widths=[3.5, 4.5, 8.5])

    out = Path(__file__).parent / "Metodologia-Narzedzia-Stack.docx"
    doc.save(out)
    print(f"✓ Zapisano: {out}")


if __name__ == "__main__":
    main()
