#!/usr/bin/env python3
"""4.4. Multi-tenant — izolacja danych na poziomie wierszy — DOCX zgodny z formatem CDV."""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run(run, *, bold=False, italic=False, size=10, font='Verdana', mono=False):
    name = 'Consolas' if mono else font
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)


def add_paragraph(doc, runs, *, indent_first=True, justify=True, space_after=8):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if indent_first:
        pf.first_line_indent = Cm(0.6)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, *style in runs:
        bold = 'b' in style
        italic = 'i' in style
        mono = 'm' in style
        r = p.add_run(text)
        set_run(r, bold=bold, italic=italic, mono=mono)
    return p


def add_heading(doc, text, *, size=11, level=2):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(8 if level == 2 else 6)
    pf.space_after = Pt(8 if level == 2 else 4)
    r = p.add_run(text)
    set_run(r, bold=True, size=size)


def add_listing(doc, *, listing_no, title, file_path, lines, comment):
    box = doc.add_paragraph()
    box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = box.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(8)
    pf.space_after = Pt(2)
    r = box.add_run(f'[ TUTAJ WSTAW ZRZUT EKRANU KODU — Listing {listing_no} ]')
    set_run(r, bold=True, italic=True)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = info.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    r1 = info.add_run('Plik: ')
    set_run(r1)
    r2 = info.add_run(file_path)
    set_run(r2, mono=True)
    r3 = info.add_run(f'   Linie: {lines}')
    set_run(r3)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = cap.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(2)
    pf.space_after = Pt(8)
    r = cap.add_run(f'Listing {listing_no}. {title}')
    set_run(r, italic=True)

    if comment:
        add_paragraph(doc, [(comment,)])


def main():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Verdana'
    style.font.size = Pt(10)

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    add_heading(doc, '4.4. Architektura multi-tenant — izolacja danych firm', size=12)

    add_paragraph(doc, [
        ('MoodFlow został zaprojektowany jako platforma ', ),
        ('multi-tenant', 'b'),
        (' (wielodostępna), co oznacza, że jedna instancja aplikacji oraz jedna baza danych obsługują wiele niezależnych firm jednocześnie. Każda firma — w terminologii systemu nazywana ', ),
        ('organizacją', 'i'),
        (' — ma własnych pracowników, własne testy, własne wyniki i własne raporty. Krytycznym wymaganiem biznesowym jest to, aby pracownik firmy A nigdy nie zobaczył danych firmy B, a administrator firmy A nie miał wglądu w listę pracowników firmy B. Naruszenie tej zasady oznaczałoby poważny incydent prywatności danych dotyczących dobrostanu psychicznego.', ),
    ])

    add_heading(doc, 'Wybór strategii izolacji danych', level=3)

    add_paragraph(doc, [
        ('Istnieją trzy klasyczne podejścia do izolacji danych w systemach multi-tenant. Pierwsze — ', ),
        ('database per tenant', 'b'),
        (' — zakłada, że każda firma ma własną osobną bazę danych. Daje to silną izolację, ale komplikuje wdrożenie i utrzymanie (każda nowa firma to osobna baza do utworzenia, migrowania i monitorowania). Drugie — ', ),
        ('schema per tenant', 'b'),
        (' — wykorzystuje mechanizm schematów PostgreSQL: jedna baza, ale każda firma ma własny zestaw tabel. To rozwiązanie pośrednie, ale wymaga dynamicznego przełączania schematu dla każdego zapytania. Trzecie — ', ),
        ('row-level isolation', 'b'),
        (' (izolacja na poziomie wierszy) — zakłada, że wszystkie firmy korzystają z tych samych tabel, ale każdy wiersz zawiera dodatkową kolumnę identyfikującą organizację, a aplikacja zawsze filtruje dane po tym identyfikatorze.', ),
    ])

    add_paragraph(doc, [
        ('W MoodFlow wybrano podejście ', ),
        ('row-level isolation', 'b'),
        (' z trzech powodów. Po pierwsze, jest najprostsze do wdrożenia — nie wymaga skomplikowanej orkiestracji baz ani schematów. Po drugie, idealnie pasuje do skali pracy inżynierskiej — system testowy nie obsługuje setek firm i nie potrzebuje fizycznej separacji storage. Po trzecie, izolacja jest realizowana w warstwie aplikacyjnej w sposób, który łatwo się audytuje i opisuje w pracy: wystarczy pokazać, że każde zapytanie do bazy zawiera warunek ', ),
        ('WHERE organizationId = :id', 'm'),
        ('. Gdyby system miał kiedyś obsługiwać kilkaset organizacji z wymogami zgodności typu HIPAA czy SOC 2, można rozważyć migrację do schema per tenant — ale dla obecnej skali byłoby to ', ),
        ('over-engineering', 'i'),
        ('.', ),
    ])

    add_heading(doc, 'Klucz organizacji w encjach', level=3)

    add_paragraph(doc, [
        ('Centralnym elementem architektury multi-tenant jest tabela ', ),
        ('organizations', 'm'),
        (' przechowująca dane firm korzystających z platformy. Każda organizacja ma unikalny identyfikator, nazwę, NIP, status cyklu życia (', ),
        ('PENDING', 'i'),
        (', ', ),
        ('ACTIVE', 'i'),
        (', ', ),
        ('BLOCKED', 'i'),
        (', ', ),
        ('REJECTED', 'i'),
        ('), kod zaproszenia oraz powiązanie z administratorem firmy. Wszystkie pozostałe tabele operacyjne — ', ),
        ('users', 'm'),
        (', ', ),
        ('departments', 'm'),
        (', ', ),
        ('assessment_assignments', 'm'),
        (', ', ),
        ('assessment_results', 'm'),
        (', ', ),
        ('audit_logs', 'm'),
        (' — zawierają kolumnę ', ),
        ('organizationId', 'b'),
        (' będącą kluczem obcym do tabeli ', ),
        ('organizations', 'm'),
        ('. To właśnie ta kolumna fizycznie identyfikuje, do której firmy należy dany rekord.', ),
    ])

    add_listing(
        doc,
        listing_no='4.13',
        title='Encja organizations — definicja firmy w systemie',
        file_path='backend-api/src/modules/organizations/entities/organization.entity.ts',
        lines='1–42',
        comment=(
            'Listing 4.13 prezentuje encję Organization. Najważniejsze pola to status '
            '(określający cykl życia firmy w systemie), inviteCode (unikalny kod zaproszenia '
            'pozwalający pracownikom rejestrować się w obrębie tej organizacji) oraz '
            'adminUserId (powiązanie z kontem administratora firmy). Dekoratory @Column '
            'TypeORM mapują pola na kolumny w PostgreSQL.'
        ),
    )

    add_paragraph(doc, [
        ('Encja ', ),
        ('User', 'm'),
        (' zawiera kolumnę ', ),
        ('organizationId', 'b'),
        (' wraz z relacją ', ),
        ('@ManyToOne', 'm'),
        (' do encji ', ),
        ('Organization', 'm'),
        ('. Dzięki temu z poziomu kodu można w łatwy sposób pobrać organizację użytkownika (', ),
        ('user.organization', 'm'),
        ('), a TypeORM automatycznie wygeneruje odpowiednie zapytanie ', ),
        ('JOIN', 'i'),
        ('. Klucz obcy ', ),
        ('organizationId', 'm'),
        (' z opcją ', ),
        ('onDelete: SET NULL', 'm'),
        (' oznacza, że jeśli organizacja zostanie usunięta z bazy, powiązani użytkownicy nie znikną automatycznie — ich pole ', ),
        ('organizationId', 'm'),
        (' zostanie wyzerowane, co umożliwia kontrolowaną migrację lub czyszczenie danych przez administratora platformy.', ),
    ])

    add_listing(
        doc,
        listing_no='4.14',
        title='Encja users z polem organizationId i relacją do organizacji',
        file_path='backend-api/src/modules/users/entities/user.entity.ts',
        lines='15–68',
        comment=(
            'Listing 4.14 pokazuje encję User. Kluczowe dla architektury multi-tenant '
            'są linie 50–55: pole organizationId oraz relacja @ManyToOne do encji '
            'Organization. Dodatkowo widoczne są pola role (employee, hr, admin, '
            'super_admin) oraz status (pending, active, suspended, rejected), '
            'które wspólnie determinują, do jakich danych użytkownik ma dostęp.'
        ),
    )

    add_heading(doc, 'Filtrowanie zapytań po organizationId', level=3)

    add_paragraph(doc, [
        ('Sama obecność kolumny ', ),
        ('organizationId', 'm'),
        (' w bazie nie wystarczy — kluczowe jest to, że ', ),
        ('każde', 'b'),
        (' zapytanie modyfikujące lub odczytujące dane musi jawnie filtrować po tej kolumnie. Konsekwentne stosowanie tej zasady realizuje serwis ', ),
        ('UsersService', 'm'),
        ('. Metoda ', ),
        ('findAll', 'm'),
        (' przyjmuje opcjonalny parametr ', ),
        ('callerOrganizationId', 'm'),
        ('. Jeśli wywołanie pochodzi od administratora firmy, parametr jest ustawiony i zapytanie zwraca tylko użytkowników z tej organizacji. Jeśli wywołanie pochodzi od właściciela platformy (super-admina), parametr jest pominięty i metoda zwraca wszystkich użytkowników. Identyczna logika obowiązuje w metodach ', ),
        ('findPending', 'm'),
        (', ', ),
        ('updateStatus', 'm'),
        (', ', ),
        ('renameDepartment', 'm'),
        (' i wszystkich innych operacjach na danych pracowników.', ),
    ])

    add_listing(
        doc,
        listing_no='4.15',
        title='Filtrowanie listy użytkowników po organizationId',
        file_path='backend-api/src/modules/users/users.service.ts',
        lines='80–103',
        comment=(
            'Listing 4.15 pokazuje, jak prosta i deklaratywna jest izolacja danych. '
            'Jeśli wywołujący ma przypisaną organizację, metoda automatycznie '
            'dodaje warunek where: { organizationId: callerOrganizationId } do '
            'zapytania TypeORM. Super-admin pomija ten warunek i widzi wszystkie '
            'organizacje. Identyczny wzorzec jest zastosowany w pozostałych metodach '
            'serwisu — dzięki temu izolacja jest ujednolicona w całej aplikacji.'
        ),
    )

    add_paragraph(doc, [
        ('Jeszcze ważniejsze są ', ),
        ('operacje modyfikujące', 'b'),
        (' — w szczególności ', ),
        ('updateStatus', 'm'),
        (' (zatwierdzanie kont) oraz ', ),
        ('renameDepartment', 'm'),
        (' (zmiana nazwy działu). Te metody dodatkowo ', ),
        ('weryfikują', 'b'),
        (' przed wykonaniem operacji, czy obiekt docelowy faktycznie należy do organizacji wywołującego, a w razie niezgodności rzucają wyjątek ', ),
        ('ForbiddenException', 'm'),
        ('. Dzięki temu nawet teoretyczna manipulacja parametrami żądania (na przykład próba zatwierdzenia konta z innej firmy przez podstawienie identyfikatora) zakończy się błędem 403, a nie wykonaniem niedozwolonej akcji.', ),
    ])

    add_listing(
        doc,
        listing_no='4.16',
        title='Weryfikacja właściciela rekordu przed operacją modyfikującą',
        file_path='backend-api/src/modules/users/users.service.ts',
        lines='106–115',
        comment=(
            'Listing 4.16 pokazuje fragment metody updateStatus. Po pobraniu '
            'użytkownika z bazy sprawdzane jest, czy jego organizationId zgadza '
            'się z organizationId wywołującego. Jeżeli nie, wyrzucany jest '
            'wyjątek ForbiddenException, a operacja w ogóle się nie wykonuje. '
            'Ta dwustopniowa weryfikacja (filtr w zapytaniu + sprawdzenie po '
            'pobraniu) realizuje zasadę defense in depth na poziomie warstwy '
            'biznesowej.'
        ),
    )

    add_heading(doc, 'Kod zaproszenia — przypisywanie pracownika do firmy', level=3)

    add_paragraph(doc, [
        ('Pracownik nie loguje się „do platformy MoodFlow" — loguje się do swojej firmy. Aby to zrealizować, przy zakładaniu nowej organizacji generowany jest unikalny ', ),
        ('kod zaproszenia', 'b'),
        (' w formacie ', ),
        ('MOOD-XXXXXXXX', 'm'),
        (' (osiem znaków heksadecymalnych zaczerpniętych z generatora kryptograficznego ', ),
        ('crypto.randomBytes', 'm'),
        ('). Administrator firmy przekazuje ten kod swoim pracownikom, którzy podają go w formularzu rejestracji. Backend wyszukuje organizację po ', ),
        ('inviteCode', 'm'),
        (' i automatycznie ustawia ', ),
        ('organizationId', 'm'),
        (' nowego użytkownika. Dzięki temu pracownik nigdy nie wybiera ręcznie firmy z listy — co eliminuje ryzyko pomyłki oraz scenariusz, w którym ktoś niezwiązany z firmą zarejestruje się w niej „przez przypadek".', ),
    ])

    add_listing(
        doc,
        listing_no='4.17',
        title='Rejestracja pracownika z walidacją kodu zaproszenia',
        file_path='backend-api/src/modules/auth/auth.service.ts',
        lines='92–119',
        comment=(
            'Listing 4.17 pokazuje metodę registerEmployee. Najpierw sprawdzane jest, '
            'czy kod zaproszenia istnieje w bazie i czy organizacja jest aktywna — '
            'jeśli nie, rejestracja kończy się błędem. Dopiero po pomyślnej walidacji '
            'tworzone jest konto pracownika z automatycznie ustawionym organizationId '
            'pobranym z odnalezionej organizacji. Status nowego konta to PENDING '
            '— pracownik czeka na zatwierdzenie przez administratora swojej firmy.'
        ),
    )

    add_heading(doc, 'Cykl życia organizacji', level=3)

    add_paragraph(doc, [
        ('Każda nowa firma trafiająca do systemu przechodzi przez kontrolowany cykl życia. Po wypełnieniu formularza rejestracyjnego organizacja otrzymuje status ', ),
        ('PENDING', 'b'),
        (' i czeka na akceptację przez właściciela platformy. Dopóki status nie zostanie zmieniony na ', ),
        ('ACTIVE', 'b'),
        (', kod zaproszenia jest nieaktywny — nie można dołączyć pracowników, nie można korzystać z aplikacji. Po akceptacji firma staje się aktywna i może rozpocząć normalną pracę. W razie naruszenia regulaminu lub zakończenia umowy właściciel platformy może ustawić status na ', ),
        ('BLOCKED', 'b'),
        (' — wszyscy użytkownicy tej firmy zostają wówczas wylogowani, a ich kolejne próby logowania kończą się komunikatem o zablokowanym koncie firmowym. Status ', ),
        ('REJECTED', 'b'),
        (' jest stosowany dla zgłoszeń odrzuconych już na etapie weryfikacji.', ),
    ])

    add_heading(doc, 'Podsumowanie izolacji multi-tenant', level=3)

    add_paragraph(doc, [
        ('Architektura multi-tenant w MoodFlow opiera się na trzech filarach: ', ),
        ('strukturalnym', 'b'),
        (' (kolumna organizationId we wszystkich tabelach operacyjnych), ', ),
        ('aplikacyjnym', 'b'),
        (' (każde zapytanie i każda mutacja filtruje po organizationId pobranym z tokena JWT) oraz ', ),
        ('proceduralnym', 'b'),
        (' (kontrolowany cykl życia organizacji od PENDING po BLOCKED). Razem te trzy warstwy zapewniają, że dane jednej firmy nie wyciekną do drugiej, a system pozostaje prosty w utrzymaniu i opisaniu w pracy inżynierskiej. Decyzję o wyborze tej strategii — zamiast bardziej złożonych alternatyw typu schema per tenant czy database per tenant — uzasadnia skala projektu: pracownia inżynierska wymaga rozwiązania działającego, prostego i jasno opisywalnego, a nie infrastruktury enterprise dla tysięcy klientów.', ),
    ])

    out = '/Users/zelek/MoodFlow/praca-pdf/4.4-Multi-tenant-izolacja-danych.docx'
    doc.save(out)
    print(f'OK: {out}')


if __name__ == '__main__':
    main()
