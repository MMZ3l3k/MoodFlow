#!/usr/bin/env python3
"""Generuje plik DOCX z 26 scenariuszami przypadków użycia w formacie identycznym
ze wzorcem z pracy inżynierskiej (Verdana 10pt, line-spacing 1.15, A4)."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tcPr.append(tcBorders)


def set_run(run, *, bold=False, italic=False, size=10, font='Verdana'):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font)
    rFonts.set(qn('w:hAnsi'), font)
    rFonts.set(qn('w:cs'), font)


def add_paragraph_in_cell(cell, text, *, bold=False, italic=False, indent_left=None,
                          line_spacing=1.15, first=False, ordered_idx=None):
    if first and cell.paragraphs and not cell.paragraphs[0].text:
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent_left is not None:
        pf.left_indent = Cm(indent_left)
    if ordered_idx is not None:
        run_n = p.add_run(f'{ordered_idx}. ')
        set_run(run_n)
    run = p.add_run(text)
    set_run(run, bold=bold, italic=italic)
    return p


def set_row_two_cells(table, label, value_lines):
    row = table.add_row()
    c1, c2 = row.cells
    c1.width = Cm(5.5)
    c2.width = Cm(11)
    set_cell_border(c1)
    set_cell_border(c2)
    c1.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    c2.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    add_paragraph_in_cell(c1, label, bold=True, first=True)

    if isinstance(value_lines, str):
        add_paragraph_in_cell(c2, value_lines, first=True)
    else:
        for i, line in enumerate(value_lines):
            if isinstance(line, tuple):
                idx, txt = line
                add_paragraph_in_cell(c2, txt, indent_left=0.6, ordered_idx=idx,
                                       first=(i == 0))
            else:
                add_paragraph_in_cell(c2, line, first=(i == 0))


def add_use_case(doc, *, lp, code, title, actors, preconditions, description,
                 main_path, alt_path, postconditions, table_no):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    set_row_two_cells(table, 'Lp.', str(lp))
    set_row_two_cells(table, 'Nazwa:', f'{code} {title}')
    set_row_two_cells(table, 'Aktorzy:', actors)
    set_row_two_cells(table, 'Warunki początkowe:', preconditions)
    set_row_two_cells(table, 'Opis:', description)
    set_row_two_cells(table, 'Ścieżka główna:',
                      [(i + 1, t) for i, t in enumerate(main_path)])
    set_row_two_cells(table, 'Ścieżka alternatywna:',
                      [(label, t) for label, t in alt_path])
    set_row_two_cells(table, 'Warunki końcowe:', postconditions)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = cap.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(2)
    pf.space_after = Pt(8)
    r = cap.add_run(f'Tabela 2.{table_no} {code} — {title}')
    set_run(r, italic=True)


def main():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Verdana'
    style.font.size = Pt(10)

    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    h = doc.add_paragraph()
    h.paragraph_format.space_after = Pt(8)
    rh = h.add_run('2.4. Scenariusze przypadków użycia')
    set_run(rh, bold=True, size=11)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.first_line_indent = Cm(0.6)
    intro.paragraph_format.line_spacing = 1.15
    intro.paragraph_format.space_after = Pt(8)
    r = intro.add_run(
        'Niniejszy podrozdział zawiera szczegółowe scenariusze przypadków użycia '
        '(oznaczonych dalej skrótem PU) zidentyfikowanych w aplikacji MoodFlow. '
        'Każdy scenariusz opisano w jednolitym formacie tabelarycznym obejmującym: '
        'nazwę przypadku użycia, listę aktorów, warunki początkowe, opis celu, '
        'ścieżkę główną, ścieżkę alternatywną oraz warunki końcowe. Łącznie '
        'wyodrębniono 26 przypadków użycia, pogrupowanych zgodnie z czterema '
        'rolami systemu: pracownik, HR, administrator firmy oraz właściciel '
        'platformy (super-administrator).')
    set_run(r)

    use_cases = [
        # PU-1
        dict(lp=1, code='PU-1', title='Logowanie pracownika',
             actors='Pracownik',
             preconditions='Pracownik posiada aktywne konto, którego status to active, oraz znajduje się na stronie logowania panelu pracownika.',
             description='Aktor uwierzytelnia się w systemie, aby uzyskać dostęp do funkcji panelu pracownika.',
             main_path=[
                 'Aktor wprowadza adres e-mail i hasło.',
                 'Aktor potwierdza dane przyciskiem „Zaloguj się”.',
                 'System weryfikuje poświadczenia (porównanie skrótu bcrypt) oraz status konta.',
                 'System wystawia tokeny dostępu i odświeżania jako bezpieczne ciasteczka HttpOnly.',
                 'Aktor zostaje przekierowany do pulpitu (dashboard) panelu pracownika.',
             ],
             alt_path=[
                 ('3a.', 'Niepoprawne dane uwierzytelniające — system wyświetla komunikat „Nieprawidłowy e-mail lub hasło”.'),
                 ('3b.', 'Konto o statusie pending_approval — system wyświetla komunikat „Konto oczekuje na zatwierdzenie przez administratora”.'),
                 ('3c.', 'Po pięciu nieudanych próbach logowania aktywuje się ograniczenie liczby żądań (rate limiting) na 5 minut.'),
             ],
             postconditions='Aktor jest zalogowany w systemie, sesja jest aktywna, a tokeny zostały zapisane w ciasteczkach przeglądarki.'),
        # PU-2
        dict(lp=2, code='PU-2', title='Rejestracja konta pracownika przez kod zaproszenia',
             actors='Pracownik (gość)',
             preconditions='Aktor otrzymał od administratora swojej firmy unikalny kod zaproszenia (inviteCode) oraz znajduje się na stronie rejestracji.',
             description='Aktor zakłada konto pracownicze w ramach istniejącej organizacji w trybie multi-tenant.',
             main_path=[
                 'Aktor wprowadza imię, nazwisko, służbowy adres e-mail oraz hasło.',
                 'Aktor wprowadza kod zaproszenia firmy.',
                 'Aktor zatwierdza zgody (regulamin, polityka prywatności RODO).',
                 'System weryfikuje poprawność kodu zaproszenia i przypisuje konto do organizacji.',
                 'System hashuje hasło algorytmem bcrypt (12 rund) i tworzy konto o statusie pending_approval.',
                 'System wysyła e-mail z informacją o oczekiwaniu na akceptację.',
             ],
             alt_path=[
                 ('4a.', 'Niepoprawny kod zaproszenia — system wyświetla komunikat „Kod zaproszenia jest nieprawidłowy lub wygasł”.'),
                 ('4b.', 'Adres e-mail jest już zarejestrowany — system wyświetla komunikat „Konto z tym adresem e-mail już istnieje”.'),
                 ('1a.', 'Hasło nie spełnia wymogów (min. 8 znaków, mieszane znaki) — system blokuje rejestrację i wskazuje brakujące kryteria.'),
             ],
             postconditions='Konto pracownika zostało utworzone w bazie ze statusem pending_approval i czeka na zatwierdzenie przez administratora firmy.'),
        # PU-3
        dict(lp=3, code='PU-3', title='Wypełnienie testu psychologicznego',
             actors='Pracownik',
             preconditions='Aktor jest zalogowany, posiada przypisany aktywny test (PHQ-9, GAD-7, PSS-10, WHO-5 lub MOOD10) oraz znajduje się w widoku „Testy”.',
             description='Aktor wypełnia kwestionariusz, którego wynik jest następnie obliczany i zapisywany przez backend.',
             main_path=[
                 'Aktor wybiera test z listy dostępnych.',
                 'System pobiera definicję testu wraz z pytaniami i opcjami odpowiedzi w skali Likerta.',
                 'Aktor odpowiada na kolejne pytania, zaznaczając jedną opcję dla każdego.',
                 'Aktor zatwierdza odpowiedzi przyciskiem „Zakończ test”.',
                 'System waliduje kompletność odpowiedzi (wszystkie pytania udzielone).',
                 'System oblicza wynik surowy oraz znormalizowany (0–100), klasyfikuje poziom nasilenia i zapisuje wynik wraz ze snapshotem odpowiedzi.',
                 'System wyświetla potwierdzenie wraz z interpretacją wyniku osobistego.',
             ],
             alt_path=[
                 ('5a.', 'Brakuje odpowiedzi na co najmniej jedno pytanie — system wyróżnia brakujące pytania i blokuje wysłanie.'),
                 ('3a.', 'Aktor przerywa wypełnianie i opuszcza widok — postęp nie jest zachowywany; aktor zaczyna test od początku przy ponownym wejściu.'),
                 ('2a.', 'Test został już wypełniony w bieżącym oknie czasowym — system wyświetla komunikat „Ten test wypełniono już w bieżącym cyklu”.'),
             ],
             postconditions='Wynik testu został zapisany w tabeli assessment_results, a odpowiedzi szczegółowe w tabeli user_responses; wynik jest dostępny w historii pracownika.'),
        # PU-4
        dict(lp=4, code='PU-4', title='Przegląd własnych wyników',
             actors='Pracownik',
             preconditions='Aktor jest zalogowany i wypełnił co najmniej jeden test psychologiczny.',
             description='Aktor przegląda historyczne wyniki testów oraz indeks dobrostanu (Wellbeing Index) wraz z trendami w czasie.',
             main_path=[
                 'Aktor wybiera widok „Wyniki” w menu nawigacyjnym.',
                 'System pobiera wyniki aktora z bazy danych z filtrem po userId.',
                 'System oblicza zagregowany Wellbeing Index (średnia ważona pięciu testów).',
                 'Aktor przegląda kafelki z wynikami testów oraz wykresy trendu.',
                 'Aktor może rozwinąć szczegóły każdego wyniku (data, wynik znormalizowany, interpretacja).',
             ],
             alt_path=[
                 ('2a.', 'Brak historii — system wyświetla komunikat „Nie wypełniono jeszcze żadnego testu” oraz proponuje przejście do widoku „Testy”.'),
             ],
             postconditions='Aktor uzyskuje wgląd w swoją historię wyników i indeks dobrostanu.'),
        # PU-5
        dict(lp=5, code='PU-5', title='Otrzymanie powiadomienia o nowym teście',
             actors='Pracownik, System mailowy SMTP',
             preconditions='Administrator firmy lub HR przypisał aktorowi nowy test (PU-17).',
             description='Aktor otrzymuje powiadomienie w aplikacji oraz wiadomość e-mail informującą o nowym teście do wypełnienia.',
             main_path=[
                 'System wykrywa zdarzenie przypisania testu (assessment_assigned).',
                 'System tworzy wpis w tabeli notifications z linkiem głębokim do testu.',
                 'System przekazuje wiadomość do System mailowy SMTP wraz z personalizowaną treścią.',
                 'Aktor po zalogowaniu widzi czerwony znacznik liczby powiadomień przy ikonie dzwonka.',
                 'Aktor otwiera listę powiadomień i klika powiadomienie.',
                 'System przekierowuje aktora do widoku konkretnego testu.',
             ],
             alt_path=[
                 ('3a.', 'System mailowy SMTP niedostępny — powiadomienie w aplikacji pozostaje aktywne, e-mail jest ponawiany asynchronicznie.'),
                 ('5a.', 'Aktor oznacza powiadomienie jako przeczytane bez otwierania linku — wpis przyjmuje status read.'),
             ],
             postconditions='Aktor został poinformowany o nowym teście i ma do niego bezpośredni dostęp.'),
        # PU-6
        dict(lp=6, code='PU-6', title='Edycja profilu pracownika',
             actors='Pracownik',
             preconditions='Aktor jest zalogowany i znajduje się w widoku „Ustawienia”.',
             description='Aktor aktualizuje dane profilowe (imię, nazwisko, motyw aplikacji). Posiada punkt rozszerzenia (extension point) umożliwiający zmianę hasła (PU-7).',
             main_path=[
                 'Aktor modyfikuje pola formularza profilu.',
                 'Aktor zatwierdza zmiany przyciskiem „Zapisz”.',
                 'System waliduje dane (długości pól, format e-mail).',
                 'System aktualizuje wiersz w tabeli users oraz tworzy wpis w audit_logs.',
                 'System wyświetla komunikat „Zmiany zostały zapisane”.',
             ],
             alt_path=[
                 ('1a.', 'Aktor wybiera opcję „Zmień hasło” — uruchamia się rozszerzenie PU-7 Zmiana hasła.'),
                 ('3a.', 'Walidacja danych nieudana — system wyświetla błędy i blokuje zapis.'),
             ],
             postconditions='Profil aktora został zaktualizowany w bazie danych, zmiana została zarejestrowana w dzienniku audytu.'),
        # PU-7
        dict(lp=7, code='PU-7', title='Zmiana hasła (rozszerzenie PU-6)',
             actors='Pracownik (rozszerzenie aplikuje się też do pozostałych ról)',
             preconditions='Aktor jest zalogowany i wybrał opcję „Zmień hasło” w widoku ustawień profilu.',
             description='Aktor zmienia własne hasło dostępu do systemu. Przypadek użycia stanowi rozszerzenie («extend») PU-6.',
             main_path=[
                 'Aktor wprowadza obecne hasło oraz dwukrotnie nowe hasło.',
                 'System weryfikuje poprawność obecnego hasła (porównanie skrótu).',
                 'System waliduje siłę nowego hasła oraz zgodność powtórzenia.',
                 'System hashuje nowe hasło (bcrypt 12 rund) i aktualizuje pole passwordHash.',
                 'System unieważnia wszystkie aktywne tokeny odświeżania (rotacja sesji).',
                 'System wyświetla komunikat „Hasło zostało zmienione”.',
             ],
             alt_path=[
                 ('2a.', 'Niepoprawne obecne hasło — system wyświetla błąd i nie zmienia hasła.'),
                 ('3a.', 'Hasła nie są zgodne — system blokuje zapis i wskazuje pole z różnicą.'),
                 ('3b.', 'Hasło nie spełnia kryteriów siły — system wskazuje brakujące wymogi.'),
             ],
             postconditions='Hasło zostało zmienione, sesje na innych urządzeniach wygasły, zdarzenie zarejestrowano w dzienniku audytu.'),
        # PU-8
        dict(lp=8, code='PU-8', title='Wylogowanie',
             actors='Pracownik (analogicznie HR, Administrator firmy, Właściciel platformy)',
             preconditions='Aktor jest zalogowany w systemie.',
             description='Aktor kończy aktywną sesję i wraca na ekran logowania.',
             main_path=[
                 'Aktor wybiera opcję „Wyloguj się” w menu profilu.',
                 'System unieważnia tokeny dostępu i odświeżania.',
                 'System usuwa ciasteczka HttpOnly (mf_access, mf_refresh).',
                 'System przekierowuje aktora do strony logowania.',
             ],
             alt_path=[
                 ('2a.', 'Sesja wygasła wcześniej — system wykonuje wylogowanie idempotentnie i przekierowuje aktora bez błędu.'),
             ],
             postconditions='Sesja aktora została zakończona, dostęp do funkcji chronionych wymaga ponownego logowania.'),
        # PU-9
        dict(lp=9, code='PU-9', title='Logowanie pracownika HR',
             actors='HR',
             preconditions='Aktor posiada konto z rolą hr, którego status to active, i znajduje się na stronie logowania panelu administracyjnego.',
             description='Aktor uwierzytelnia się w celu uzyskania dostępu do panelu HR z funkcjami analityki anonimowej.',
             main_path=[
                 'Aktor wprowadza adres e-mail oraz hasło.',
                 'System weryfikuje poświadczenia oraz rolę użytkownika.',
                 'System wystawia tokeny i przekierowuje aktora na pulpit HR.',
             ],
             alt_path=[
                 ('2a.', 'Konto nie posiada roli hr — system zwraca komunikat „Brak uprawnień do panelu HR”.'),
                 ('2b.', 'Niepoprawne dane uwierzytelniające — system wyświetla komunikat „Nieprawidłowy e-mail lub hasło”.'),
             ],
             postconditions='Aktor jest zalogowany i ma dostęp do widoków przeznaczonych dla roli HR.'),
        # PU-10
        dict(lp=10, code='PU-10', title='Anonimowa analityka dobrostanu (k-anonimowość)',
             actors='HR',
             preconditions='Aktor jest zalogowany w panelu HR. Organizacja posiada co najmniej k=5 pracowników, którzy wypełnili przynajmniej jeden test.',
             description='Aktor analizuje zagregowane wyniki dobrostanu pracowników z gwarancją prywatności na poziomie k=5.',
             main_path=[
                 'Aktor wybiera widok „Analityka”.',
                 'System pobiera zagregowane dane z tabeli assessment_results z filtrem po organizationId.',
                 'System weryfikuje próg k-anonimowości (n ≥ 5) na poziomie organizacji oraz każdej grupy.',
                 'System renderuje wykresy: indeks dobrostanu, rozkład nasilenia, trend w czasie.',
             ],
             alt_path=[
                 ('3a.', 'Liczba uczestników poniżej progu k=5 — system wyświetla komunikat „Zbyt mało danych, aby zachować anonimowość” i ukrywa wykres.'),
                 ('2a.', 'Brak wyników w organizacji — system pokazuje stan pusty i sugeruje przypisanie testów.'),
             ],
             postconditions='Aktor uzyskuje wgląd w zagregowane wskaźniki dobrostanu bez możliwości deanonimizacji pojedynczych pracowników.'),
        # PU-11
        dict(lp=11, code='PU-11', title='Filtrowanie wyników po działach',
             actors='HR',
             preconditions='Aktor jest zalogowany w panelu HR i znajduje się w widoku analityki.',
             description='Aktor zawęża zakres analizy do wybranego działu lub grupy działów, zachowując zasady anonimizacji.',
             main_path=[
                 'Aktor otwiera filtr „Dział” i wybiera jeden lub kilka działów organizacji.',
                 'Aktor zatwierdza wybór przyciskiem „Zastosuj filtry”.',
                 'System pobiera zagregowane dane z dodatkowym filtrem po departmentId.',
                 'System ponownie weryfikuje próg k-anonimowości na poziomie wyfiltrowanej grupy.',
                 'System wyświetla zaktualizowane wykresy.',
             ],
             alt_path=[
                 ('4a.', 'Wybrany dział ma mniej niż k=5 osób — system wyświetla komunikat „Wyniki niewidoczne — zbyt mała grupa” i sugeruje rozszerzenie zakresu.'),
                 ('1a.', 'Aktor wybiera „Wszystkie działy” — filtr nie jest stosowany.'),
             ],
             postconditions='Wyniki zostały zawężone do wskazanych działów lub system poinformował o niemożności prezentacji.'),
        # PU-12
        dict(lp=12, code='PU-12', title='Eksport raportu CSV/PDF',
             actors='HR',
             preconditions='Aktor jest zalogowany i znajduje się w widoku analityki z aktywnymi filtrami spełniającymi próg k-anonimowości.',
             description='Aktor pobiera raport zagregowanych wyników w formacie CSV lub PDF.',
             main_path=[
                 'Aktor wybiera przycisk „Eksportuj” i wskazuje format (CSV lub PDF).',
                 'System przygotowuje zestaw zagregowanych danych zgodnie z aktywnymi filtrami.',
                 'System weryfikuje uprawnienia roli oraz próg k-anonimowości eksportowanych grup.',
                 'System generuje plik i odsyła go do przeglądarki aktora.',
                 'System tworzy wpis w audit_logs o eksporcie raportu.',
             ],
             alt_path=[
                 ('3a.', 'Próg k-anonimowości nie jest spełniony — system blokuje eksport i wyświetla komunikat o niemożności udostępnienia danych.'),
                 ('4a.', 'Błąd generowania pliku — system wyświetla komunikat o spróbowaniu ponownie.'),
             ],
             postconditions='Aktor pobrał plik raportu, zdarzenie zostało zarejestrowane w dzienniku audytu.'),
        # PU-13
        dict(lp=13, code='PU-13', title='Przegląd alertów dobrostanu',
             actors='HR',
             preconditions='Aktor jest zalogowany. W organizacji wystąpiły zagregowane wyniki przekraczające zdefiniowane progi alertowe (np. wzrost depresji w dziale).',
             description='Aktor przegląda automatycznie wygenerowane alerty dotyczące poziomu dobrostanu w działach.',
             main_path=[
                 'System cyklicznie analizuje wyniki i tworzy alerty dla działów spełniających próg k=5 i przekraczających próg ryzyka.',
                 'Aktor otwiera widok „Alerty”.',
                 'System pobiera listę aktywnych alertów dla organizacji.',
                 'Aktor przegląda szczegóły alertu (test, dział, zmiana w czasie).',
                 'Aktor oznacza alert jako „obsłużony” lub dodaje notatkę.',
             ],
             alt_path=[
                 ('3a.', 'Brak aktywnych alertów — system wyświetla informację „Brak alertów wymagających uwagi”.'),
             ],
             postconditions='Aktor zna aktualne sygnały ryzyka w organizacji, alerty zostały oznaczone zgodnie z reakcją.'),
        # PU-14
        dict(lp=14, code='PU-14', title='Logowanie administratora firmy',
             actors='Administrator firmy',
             preconditions='Aktor posiada konto z rolą admin, status active, organizacja jest aktywna.',
             description='Aktor uwierzytelnia się w celu zarządzania pracownikami i konfiguracją swojej organizacji.',
             main_path=[
                 'Aktor wprowadza dane uwierzytelniające w panelu administratora.',
                 'System weryfikuje rolę i status organizacji.',
                 'System przekierowuje aktora na pulpit administratora firmy.',
             ],
             alt_path=[
                 ('2a.', 'Organizacja ma status blocked — system uniemożliwia logowanie i wyświetla komunikat „Konto firmowe zostało zablokowane”.'),
                 ('2b.', 'Niepoprawne dane uwierzytelniające — system zwraca komunikat błędu.'),
             ],
             postconditions='Aktor jest zalogowany w panelu administratora firmy.'),
        # PU-15
        dict(lp=15, code='PU-15', title='Zaproszenie pracownika',
             actors='Administrator firmy, System mailowy SMTP',
             preconditions='Aktor jest zalogowany w panelu administratora firmy.',
             description='Aktor generuje kod zaproszenia i wysyła go nowemu pracownikowi w celu założenia konta.',
             main_path=[
                 'Aktor wybiera widok „Pracownicy” → „Zaproś pracownika”.',
                 'Aktor wprowadza adres e-mail kandydata oraz opcjonalnie dział.',
                 'System generuje unikalny kod zaproszenia powiązany z organizacją.',
                 'System przekazuje wiadomość do System mailowy SMTP z linkiem rejestracyjnym.',
                 'System tworzy wpis w audit_logs o wysłaniu zaproszenia.',
             ],
             alt_path=[
                 ('2a.', 'Aktor wprowadził adres e-mail już istniejącego pracownika — system wyświetla komunikat „Pracownik o tym adresie już istnieje”.'),
                 ('4a.', 'System mailowy SMTP niedostępny — system pokazuje wygenerowany kod do przekazania ręcznie.'),
             ],
             postconditions='Pracownik otrzymał kod zaproszenia drogą e-mailową lub jako kod do przekazania ręcznie.'),
        # PU-16
        dict(lp=16, code='PU-16', title='Zarządzanie pracownikami',
             actors='Administrator firmy',
             preconditions='Aktor jest zalogowany. Istnieją konta pracowników w organizacji o różnych statusach.',
             description='Aktor zatwierdza nowo zarejestrowanych pracowników, blokuje konta lub zmienia ich role w ramach organizacji.',
             main_path=[
                 'Aktor otwiera widok „Pracownicy”.',
                 'System wyświetla listę pracowników filtrowaną po organizationId.',
                 'Aktor wybiera pracownika oczekującego na zatwierdzenie i klika „Zatwierdź”.',
                 'System aktualizuje status konta na active i tworzy wpis w audit_logs.',
                 'System wysyła do pracownika powiadomienie e-mail o aktywacji konta.',
             ],
             alt_path=[
                 ('3a.', 'Aktor odrzuca konto — status zmienia się na rejected, pracownik traci możliwość logowania.'),
                 ('3b.', 'Aktor zmienia rolę pracownika (np. employee → hr) — system aktualizuje rekord po dodatkowym potwierdzeniu.'),
                 ('3c.', 'Aktor zawiesza konto — status zmienia się na suspended, sesje pracownika są unieważniane.'),
             ],
             postconditions='Status, rola lub stan konta pracownika został zmieniony zgodnie z decyzją aktora; zdarzenie zarejestrowano w dzienniku audytu.'),
        # PU-17
        dict(lp=17, code='PU-17', title='Przypisanie testu',
             actors='Administrator firmy',
             preconditions='Aktor jest zalogowany w panelu administratora firmy. W systemie zdefiniowane są aktywne szablony testów.',
             description='Aktor przypisuje wybrany test do wszystkich pracowników, do wybranego działu lub do konkretnego pracownika. Posiada punkt rozszerzenia: PU-18 (Wybór działu docelowego).',
             main_path=[
                 'Aktor wybiera widok „Testy” → „Przypisz nowy test”.',
                 'Aktor wybiera szablon testu z listy (PHQ-9, GAD-7, PSS-10, WHO-5, MOOD10).',
                 'Aktor wybiera typ celu: ALL, USER lub DEPARTMENT.',
                 'Aktor wskazuje przedział czasowy dostępności testu.',
                 'System zapisuje rekord w assessment_assignments.',
                 'System wyzwala procesy wysyłki powiadomień (PU-5) dla wszystkich docelowych pracowników.',
             ],
             alt_path=[
                 ('3a.', 'Aktor wybiera DEPARTMENT — uruchamia się rozszerzenie PU-18 Wybór działu docelowego.'),
                 ('3b.', 'Aktor wybiera USER — system pokazuje wyszukiwarkę pracowników z autouzupełnianiem.'),
                 ('4a.', 'Brak nakładania się okna dostępności z aktualną datą — system zezwala na zaplanowanie testu w przyszłości.'),
             ],
             postconditions='Test został przypisany do określonej grupy odbiorców, powiadomienia zostały zakolejkowane.'),
        # PU-18
        dict(lp=18, code='PU-18', title='Wybór działu docelowego (rozszerzenie PU-17)',
             actors='Administrator firmy',
             preconditions='Aktor wybrał typ celu DEPARTMENT w PU-17 Przypisanie testu.',
             description='Aktor wskazuje konkretny dział organizacji, którego pracownicy otrzymają test.',
             main_path=[
                 'System wyświetla listę działów organizacji.',
                 'Aktor wybiera jeden lub wiele działów.',
                 'System aktualizuje pole targetDepartment rekordu przypisania.',
                 'Sterowanie wraca do PU-17 (krok 5).',
             ],
             alt_path=[
                 ('1a.', 'Brak działów w organizacji — system proponuje przejście do PU-19 Konfiguracja działów.'),
             ],
             postconditions='Dział docelowy został wybrany; przypisanie testu zostanie zakończone w PU-17.'),
        # PU-19
        dict(lp=19, code='PU-19', title='Konfiguracja działów organizacji',
             actors='Administrator firmy',
             preconditions='Aktor jest zalogowany w panelu administratora firmy.',
             description='Aktor zarządza strukturą działów: tworzy, edytuje i usuwa działy.',
             main_path=[
                 'Aktor otwiera widok „Działy”.',
                 'System wyświetla listę działów organizacji.',
                 'Aktor dodaje nowy dział, podając nazwę.',
                 'System tworzy rekord w departments z polem organizationId.',
                 'System odświeża listę.',
             ],
             alt_path=[
                 ('3a.', 'Aktor edytuje istniejący dział — system aktualizuje nazwę po walidacji.'),
                 ('3b.', 'Aktor usuwa dział, do którego są przypisani pracownicy — system blokuje operację i wymaga przeniesienia pracowników.'),
             ],
             postconditions='Struktura działów organizacji została zaktualizowana.'),
        # PU-20
        dict(lp=20, code='PU-20', title='Przegląd zagregowanych wyników firmy',
             actors='Administrator firmy',
             preconditions='Aktor jest zalogowany. W organizacji istnieją wyniki testów spełniające próg k=5.',
             description='Aktor przegląda zagregowane wyniki dobrostanu w swojej firmie z poszanowaniem reguł anonimizacji.',
             main_path=[
                 'Aktor wybiera widok „Wyniki firmy”.',
                 'System pobiera zagregowane dane filtrowane po organizationId.',
                 'System weryfikuje próg k-anonimowości i ukrywa zbyt małe podgrupy.',
                 'System wyświetla indeks dobrostanu, rozkład poziomów ryzyka oraz uczestnictwo.',
             ],
             alt_path=[
                 ('3a.', 'Liczba pracowników z wynikami poniżej 5 — system wyświetla komunikat „Zbyt mała liczba uczestników do prezentacji wyników”.'),
             ],
             postconditions='Aktor uzyskuje obraz dobrostanu w organizacji bez naruszenia anonimowości pracowników.'),
        # PU-21
        dict(lp=21, code='PU-21', title='Logowanie właściciela platformy (super-admin)',
             actors='Właściciel platformy',
             preconditions='Aktor posiada konto z rolą super_admin i znajduje się na dedykowanej stronie logowania super-admina.',
             description='Aktor uwierzytelnia się w celu zarządzania całą platformą MoodFlow.',
             main_path=[
                 'Aktor wprowadza dane uwierzytelniające na stronie /super-admin/login.',
                 'System weryfikuje poświadczenia oraz rolę super_admin.',
                 'System przekierowuje aktora na pulpit właściciela platformy.',
             ],
             alt_path=[
                 ('2a.', 'Konto nie ma roli super_admin — system odmawia dostępu i loguje incydent.'),
                 ('2b.', 'Niepoprawne dane uwierzytelniające — system zwraca komunikat błędu.'),
             ],
             postconditions='Aktor jest zalogowany w panelu właściciela platformy.'),
        # PU-22
        dict(lp=22, code='PU-22', title='Akceptacja rejestracji firmy',
             actors='Właściciel platformy, System mailowy SMTP',
             preconditions='Aktor jest zalogowany. W systemie istnieje co najmniej jedna organizacja o statusie pending.',
             description='Aktor zatwierdza lub odrzuca zgłoszenia rejestracyjne nowych firm.',
             main_path=[
                 'Aktor otwiera widok „Firmy oczekujące”.',
                 'System wyświetla listę firm ze statusem pending.',
                 'Aktor weryfikuje dane firmy (NIP, opis, dane administratora).',
                 'Aktor klika „Zatwierdź”.',
                 'System zmienia status organizacji na active, generuje pierwszy inviteCode i tworzy wpis w audit_logs.',
                 'System przekazuje wiadomość do System mailowy SMTP z informacją do administratora firmy.',
             ],
             alt_path=[
                 ('4a.', 'Aktor odrzuca zgłoszenie — status zmienia się na rejected, administrator firmy otrzymuje uzasadnienie odmowy.'),
                 ('3a.', 'Aktor żąda dodatkowych informacji — wysyła pytanie e-mailem; status pozostaje pending.'),
             ],
             postconditions='Status organizacji zaktualizowano, administrator firmy został powiadomiony.'),
        # PU-23
        dict(lp=23, code='PU-23', title='Blokada lub aktywacja firmy',
             actors='Właściciel platformy',
             preconditions='Aktor jest zalogowany. W systemie istnieją organizacje aktywne lub zablokowane.',
             description='Aktor zmienia status istniejącej organizacji w celu zatrzymania lub przywrócenia jej dostępu do platformy.',
             main_path=[
                 'Aktor otwiera listę wszystkich organizacji.',
                 'Aktor wybiera organizację i klika „Zablokuj”.',
                 'System wyświetla okno potwierdzenia z polem na uzasadnienie.',
                 'Aktor wprowadza powód i potwierdza operację.',
                 'System ustawia status organizacji na blocked, unieważnia sesje wszystkich użytkowników firmy oraz tworzy wpis w audit_logs.',
             ],
             alt_path=[
                 ('2a.', 'Aktor odblokowuje wcześniej zablokowaną firmę — status wraca do active, użytkownicy odzyskują możliwość logowania.'),
                 ('4a.', 'Aktor anuluje operację — status pozostaje bez zmian.'),
             ],
             postconditions='Status organizacji został zmieniony, zdarzenie zarejestrowano w dzienniku audytu.'),
        # PU-24
        dict(lp=24, code='PU-24', title='Globalne statystyki platformy',
             actors='Właściciel platformy',
             preconditions='Aktor jest zalogowany jako super-admin.',
             description='Aktor analizuje wskaźniki działania całej platformy w przekroju wszystkich organizacji.',
             main_path=[
                 'Aktor wybiera widok „Statystyki globalne”.',
                 'System pobiera zagregowane metryki: liczba firm, liczba aktywnych pracowników, liczba wypełnionych testów, średnie czasy odpowiedzi.',
                 'System renderuje pulpit z kafelkami metryk i wykresami.',
                 'Aktor może filtrować dane po przedziale czasowym lub statusie organizacji.',
             ],
             alt_path=[
                 ('2a.', 'Brak danych w wybranym przedziale — system wyświetla stan pusty z sugestią rozszerzenia zakresu.'),
             ],
             postconditions='Aktor uzyskuje wgląd w stan techniczno-biznesowy całej platformy MoodFlow.'),
        # PU-25
        dict(lp=25, code='PU-25', title='Zarządzanie szablonami testów psychologicznych',
             actors='Właściciel platformy',
             preconditions='Aktor jest zalogowany jako super-admin.',
             description='Aktor zarządza biblioteką szablonów testów psychologicznych dostępnych w platformie (PHQ-9, GAD-7, PSS-10, WHO-5, MOOD10) oraz ich aktywnością.',
             main_path=[
                 'Aktor otwiera widok „Szablony testów”.',
                 'System wyświetla listę szablonów wraz z ich kodem, wersją, liczbą pytań i flagą isActive.',
                 'Aktor edytuje wybrany szablon: nazwę, opis, ramy czasowe, anonimizację dla HR.',
                 'System zapisuje zmiany i aktualizuje pole version szablonu.',
             ],
             alt_path=[
                 ('3a.', 'Aktor dezaktywuje szablon — flagę isActive ustawiono na false; dotychczasowe wyniki pozostają w systemie.'),
                 ('3b.', 'Aktor reaktywuje szablon — flaga wraca na true, szablon ponownie pojawia się w panelu administratora firmy.'),
             ],
             postconditions='Biblioteka szablonów testów została zaktualizowana zgodnie z decyzjami aktora.'),
        # PU-26
        dict(lp=26, code='PU-26', title='Przegląd globalnego dziennika audytu',
             actors='Właściciel platformy',
             preconditions='Aktor jest zalogowany jako super-admin.',
             description='Aktor przegląda dziennik wszystkich krytycznych zdarzeń systemowych w celach bezpieczeństwa i zgodności z RODO.',
             main_path=[
                 'Aktor otwiera widok „Dziennik audytu”.',
                 'System pobiera wpisy z tabeli audit_logs z domyślnym sortowaniem malejąco po dacie.',
                 'Aktor stosuje filtry (akcja, organizacja, użytkownik, zakres czasowy).',
                 'System aktualizuje listę wpisów zgodnie z filtrami.',
                 'Aktor może rozwinąć szczegóły zdarzenia (metadane jsonb, adres IP, użytkownik wykonujący).',
             ],
             alt_path=[
                 ('3a.', 'Brak wpisów spełniających filtry — system wyświetla informację „Brak wyników dla wybranych kryteriów”.'),
                 ('4a.', 'Aktor eksportuje wyniki do pliku CSV — operacja sama jest logowana w dzienniku audytu.'),
             ],
             postconditions='Aktor uzyskał wgląd w wybrane zdarzenia systemowe; ewentualny eksport został zarejestrowany.'),
    ]

    for i, uc in enumerate(use_cases, start=1):
        add_use_case(doc, table_no=i, **uc)

    out = '/Users/zelek/MoodFlow/praca-pdf/2.4-Scenariusze-przypadkow-uzycia.docx'
    doc.save(out)
    print(f'OK: {out}')


if __name__ == '__main__':
    main()
